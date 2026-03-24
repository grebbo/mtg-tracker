"""Import historical data from the original tracking Excel file."""
from datetime import datetime

from openpyxl import load_workbook

from app import db
from app.models import Commander, Game, GameEntry, Player

_KNOWN_PILOTS = {'enrico', 'daniele', 'lorenzo', 'simone', 'semir', 'federico'}


def _get_or_create_player(name: str) -> Player:
    p = Player.query.filter_by(name=name).first()
    if not p:
        p = Player(name=name)
        db.session.add(p)
        db.session.flush()
    return p


def _get_or_create_commander(name: str, color: str | None = None) -> Commander:
    c = Commander.query.filter_by(name=name).first()
    if not c:
        c = Commander(name=name, color_identity=color or '?')
        db.session.add(c)
        db.session.flush()
    elif color and c.color_identity == '?':
        c.color_identity = color
    return c


def _parse_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    s = str(val).strip()
    for fmt in ('%m/%d/%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _safe_int(val):
    if val is None:
        return None
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return None


def _find_header_row(ws):
    """Find the row index (1-based) of the header, skipping empty rows at top."""
    for row in ws.iter_rows(min_row=1, max_row=5):
        values = [c.value for c in row if c.value is not None]
        if values:
            return row[0].row
    return 1


def _detect_columns(header_row_values: list) -> dict:
    """Map field names to column indices from the header row."""
    defaults = {
        'date': 0, 'deck': 1, 'pilot': 2, 'colour': 3,
        'finish': 4, 'turn_order': 5, 'sol_ring': 6,
        'turns': 7, 'turn_eliminated': 8, 'eliminated_by': 9,
        'victory_condition': 10, 'loss_condition': 11, 'notes': 12,
    }
    keywords = {
        'date':              ['date'],
        'deck':              ['deck', 'archtype', 'commander'],
        'pilot':             ['pilot', 'player'],
        'colour':            ['colour', 'color'],
        'finish':            ['finish'],
        'turn_order':        ['turn order', 'order'],
        'sol_ring':          ['sol ring'],
        'turns':             ['turns'],
        'turn_eliminated':   ['turn elim'],
        'eliminated_by':     ['eliminated by'],
        'victory_condition': ['victory'],
        'loss_condition':    ['loss'],
        'notes':             ['notes'],
    }
    col_map = dict(defaults)
    for i, val in enumerate(header_row_values):
        if not val:
            continue
        cell = str(val).lower().strip()
        for field, kws in keywords.items():
            if any(kw in cell for kw in kws):
                col_map[field] = i
                break
    return col_map


def _cell(row, idx):
    try:
        return row[idx]
    except IndexError:
        return None


def import_from_excel(filepath: str) -> int:
    wb = load_workbook(filepath, data_only=True)

    # Find the tracking/game data sheet
    ws = wb.worksheets[0]
    for sheet in wb.worksheets:
        name = sheet.title.lower()
        if 'game data' in name or 'track' in name:
            ws = sheet
            break

    # Find header row (skips leading empty rows)
    header_row_num = _find_header_row(ws)
    header_values = [c.value for c in ws[header_row_num]]
    cols = _detect_columns(header_values)

    games_imported = 0
    current_date = None
    current_turns = None
    current_rows: list[dict] = []

    def flush():
        nonlocal games_imported
        if not current_rows or not current_date:
            return

        game = Game(date=current_date, total_turns=current_turns)
        db.session.add(game)
        db.session.flush()

        for rd in current_rows:
            if not rd.get('deck') or not rd.get('pilot'):
                continue

            player    = _get_or_create_player(rd['pilot'])
            commander = _get_or_create_commander(rd['deck'], rd.get('colour'))

            elim_by_id = None
            if rd.get('eliminated_by'):
                elim_cmd   = _get_or_create_commander(rd['eliminated_by'])
                elim_by_id = elim_cmd.id

            is_winner = rd.get('finish') == 1
            entry = GameEntry(
                game_id=game.id,
                player_id=player.id,
                commander_id=commander.id,
                finish=rd.get('finish'),
                turn_order=rd.get('turn_order'),
                sol_ring_t1=rd.get('sol_ring', False),
                turn_eliminated=None if is_winner else rd.get('turn_eliminated'),
                eliminated_by_id=None if is_winner else elim_by_id,
                victory_condition=rd.get('victory_condition') if is_winner else None,
                loss_condition=rd.get('loss_condition') if not is_winner else None,
                notes=rd.get('notes'),
            )
            db.session.add(entry)

        db.session.commit()
        games_imported += 1

    # Iterate data rows (start after header)
    for row in ws.iter_rows(min_row=header_row_num + 1, values_only=True):
        date_val = _cell(row, cols['date'])
        deck     = str(_cell(row, cols['deck'])).strip()  if _cell(row, cols['deck'])  else None
        pilot    = str(_cell(row, cols['pilot'])).strip() if _cell(row, cols['pilot']) else None

        # Skip completely empty rows
        if not deck and not pilot and not date_val:
            continue

        # Skip rows where pilot is not a known player (reference/stats rows)
        if pilot and pilot.lower() not in _KNOWN_PILOTS:
            continue

        colour       = str(_cell(row, cols['colour'])).strip().upper() if _cell(row, cols['colour']) else None
        finish       = _safe_int(_cell(row, cols['finish']))
        turn_order   = _safe_int(_cell(row, cols['turn_order']))
        sol_ring_raw = _cell(row, cols['sol_ring'])
        sol_ring     = str(sol_ring_raw).strip().lower() in ('yes', 'true', '1') if sol_ring_raw else False
        turns        = _safe_int(_cell(row, cols['turns']))
        turn_elim    = _safe_int(_cell(row, cols['turn_eliminated']))
        elim_by      = str(_cell(row, cols['eliminated_by'])).strip() if _cell(row, cols['eliminated_by']) else None
        victory_cond = str(_cell(row, cols['victory_condition'])).strip() if _cell(row, cols['victory_condition']) else None
        loss_cond    = str(_cell(row, cols['loss_condition'])).strip()    if _cell(row, cols['loss_condition'])    else None
        notes        = str(_cell(row, cols['notes'])).strip()             if _cell(row, cols['notes'])             else None

        # New date in col[0] = new game
        if date_val:
            parsed = _parse_date(date_val)
            if parsed:
                if current_rows:
                    flush()
                    current_rows = []
                    current_turns = None
                current_date = parsed

        if turns:
            current_turns = turns

        if deck and pilot:
            current_rows.append({
                'deck': deck, 'pilot': pilot, 'colour': colour,
                'finish': finish, 'turn_order': turn_order, 'sol_ring': sol_ring,
                'turn_eliminated': turn_elim, 'eliminated_by': elim_by,
                'victory_condition': victory_cond, 'loss_condition': loss_cond,
                'notes': notes,
            })

    if current_rows:
        flush()

    return games_imported
