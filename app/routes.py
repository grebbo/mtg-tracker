import json
import statistics
import os
from collections import defaultdict
from datetime import date, datetime

from flask import (Blueprint, current_app, flash, jsonify, redirect,
                   render_template, request, send_file, url_for)

from app import db
from app.models import (LOSS_CONDITIONS, WIN_CONDITIONS, Commander, Game,
                        GameEntry, Player)

main = Blueprint('main', __name__)


@main.route('/debug/excel', methods=['GET', 'POST'])
def debug_excel():
    """Diagnostic: show raw Excel rows to verify column positions."""
    result = None
    if request.method == 'POST':
        f = request.files.get('file')
        if f and f.filename.endswith('.xlsx'):
            import tempfile
            from openpyxl import load_workbook
            from app.import_excel import _detect_columns
            with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                f.save(tmp.name)
                tmp_path = tmp.name
            try:
                wb = load_workbook(tmp_path, data_only=True)
                ws = wb.worksheets[0]
                for sheet in wb.worksheets:
                    if 'track' in sheet.title.lower():
                        ws = sheet
                        break
                header = [c.value for c in ws[1]]
                cols = _detect_columns(header)
                rows = []
                for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True)):
                    rows.append({'index': i, 'values': list(row[:15])})
                result = {'sheet': ws.title, 'header': header[:15],
                          'cols': cols, 'rows': rows}
            finally:
                os.unlink(tmp_path)
    return render_template('debug_excel.html', result=result)


# ── helpers ──────────────────────────────────────────────────────────────────

def _compute_player_stats():
    players = Player.query.order_by(Player.name).all()
    result = []
    for p in players:
        entries = GameEntry.query.filter_by(player_id=p.id).all()
        gp = len(entries)
        wins = sum(1 for e in entries if e.finish == 1)
        result.append({
            'name': p.name,
            'games': gp,
            'wins': wins,
            'win_rate': round(wins / gp * 100, 1) if gp else 0,
            'sol_rings': sum(1 for e in entries if e.sol_ring_t1),
        })
    return result


def _compute_commander_stats():
    cmds = Commander.query.order_by(Commander.name).all()
    result = []
    for c in cmds:
        entries = GameEntry.query.filter_by(commander_id=c.id).all()
        gp = len(entries)
        if gp == 0:
            continue
        wins = sum(1 for e in entries if e.finish == 1)
        kills = GameEntry.query.filter_by(eliminated_by_id=c.id).count()
        result.append({
            'id': c.id,
            'name': c.name,
            'color': c.color_identity,
            'games': gp,
            'wins': wins,
            'kills': kills,
            'win_rate': round(wins / gp * 100, 1) if gp else 0,
        })
    result.sort(key=lambda x: x['games'], reverse=True)
    return result


def _compute_color_stats():
    cmds = Commander.query.all()
    color_data = defaultdict(lambda: {'games': 0, 'wins': 0})
    for c in cmds:
        entries = GameEntry.query.filter_by(commander_id=c.id).all()
        if not entries:
            continue
        color_data[c.color_identity]['games'] += len(entries)
        color_data[c.color_identity]['wins'] += sum(1 for e in entries if e.finish == 1)
    result = []
    for color, d in sorted(color_data.items()):
        gp = d['games']
        w = d['wins']
        result.append({
            'color': color,
            'games': gp,
            'wins': w,
            'losses': gp - w,
            'win_rate': round(w / gp * 100, 1) if gp else 0,
        })
    return result


def _compute_aggregated_color_stats():
    """For each single color (W/U/B/R/G), aggregate games/wins from ALL
    commanders whose color identity contains that letter."""
    SINGLE = ['W', 'U', 'B', 'R', 'G']
    agg = {c: {'games': 0, 'wins': 0} for c in SINGLE}
    for cmd in Commander.query.all():
        entries = GameEntry.query.filter_by(commander_id=cmd.id).all()
        if not entries:
            continue
        gp = len(entries)
        w  = sum(1 for e in entries if e.finish == 1)
        for ch in cmd.color_identity:
            if ch in agg:
                agg[ch]['games'] += gp
                agg[ch]['wins']  += w
    result = []
    for ch in SINGLE:
        gp = agg[ch]['games']
        w  = agg[ch]['wins']
        result.append({
            'color':    ch,
            'games':    gp,
            'wins':     w,
            'losses':   gp - w,
            'win_rate': round(w / gp * 100, 1) if gp else 0,
        })
    return result


def _sync_excel():
    try:
        from app.excel_sync import sync_to_excel
        sync_to_excel(current_app.config['DATA_DIR'])
    except Exception as e:
        flash(f'Game saved but Excel sync failed: {e}', 'warning')


# ── dashboard ─────────────────────────────────────────────────────────────────

@main.route('/')
def dashboard():
    total_games = Game.query.count()
    player_stats = _compute_player_stats()
    commander_stats = _compute_commander_stats()
    color_stats      = _compute_color_stats()
    agg_color_stats  = _compute_aggregated_color_stats()

    def _game_winner_info(game):
        entry = next((e for e in game.entries if e.finish == 1), None)
        if not entry:
            return None
        return {
            'player':    entry.player.name,
            'commander': entry.commander.name if entry.commander else '—',
        }

    games_with_turns = [g for g in Game.query.all() if g.total_turns]
    lengths = [g.total_turns for g in games_with_turns]
    if lengths:
        max_game = max(games_with_turns, key=lambda g: g.total_turns)
        min_game = min(games_with_turns, key=lambda g: g.total_turns)
        length_stats = {
            'avg':        round(statistics.mean(lengths), 1),
            'max':        max(lengths),
            'min':        min(lengths),
            'median':     round(statistics.median(lengths), 1),
            'max_winner': _game_winner_info(max_game),
            'min_winner': _game_winner_info(min_game),
        }
    else:
        length_stats = {
            'avg': 0, 'max': 0, 'min': 0, 'median': 0,
            'max_winner': None, 'min_winner': None,
        }

    sorted_players = sorted(player_stats, key=lambda p: p['wins'], reverse=True)
    runner_up_player = sorted_players[1] if len(sorted_players) > 1 else None

    sorted_cmds = sorted(commander_stats, key=lambda c: c['win_rate'], reverse=True)
    runner_up_cmd = sorted_cmds[1] if len(sorted_cmds) > 1 else None

    recent_games = Game.query.order_by(Game.date.desc()).limit(5).all()

    return render_template(
        'dashboard.html',
        total_games=total_games,
        player_stats=player_stats,
        commander_stats=commander_stats,
        color_stats=color_stats,
        length_stats=length_stats,
        recent_games=recent_games,
        runner_up_player=runner_up_player,
        runner_up_cmd=runner_up_cmd,
        player_stats_json=json.dumps(player_stats),
        commander_stats_json=json.dumps(commander_stats[:15]),
        color_stats_json=json.dumps(color_stats),
        agg_color_stats_json=json.dumps(agg_color_stats),
    )


# ── add / edit game ───────────────────────────────────────────────────────────

@main.route('/game/new', methods=['GET', 'POST'])
def game_new():
    players = Player.query.order_by(Player.name).all()
    commanders = Commander.query.order_by(Commander.name).all()

    if request.method == 'POST':
        try:
            game_date = datetime.strptime(request.form['date'], '%Y-%m-%d').date()
            total_turns = request.form.get('total_turns', type=int)
            num_players = request.form.get('num_players', type=int, default=4)

            game = Game(date=game_date, total_turns=total_turns)
            db.session.add(game)
            db.session.flush()

            for i in range(num_players):
                player_id = request.form.get(f'p{i}_player', type=int)
                commander_id = request.form.get(f'p{i}_commander', type=int)
                finish = request.form.get(f'p{i}_finish', type=int)
                turn_order = request.form.get(f'p{i}_turn_order', type=int)
                sol_ring = bool(request.form.get(f'p{i}_sol_ring'))
                turn_elim = request.form.get(f'p{i}_turn_eliminated', type=int)
                elim_by = request.form.get(f'p{i}_eliminated_by', type=int)
                vcond = request.form.get(f'p{i}_victory_condition', '').strip() or None
                lcond = request.form.get(f'p{i}_loss_condition', '').strip() or None
                notes = request.form.get(f'p{i}_notes', '').strip() or None

                is_winner = (finish == 1)
                entry = GameEntry(
                    game_id=game.id,
                    player_id=player_id,
                    commander_id=commander_id,
                    finish=finish,
                    turn_order=turn_order,
                    sol_ring_t1=sol_ring,
                    turn_eliminated=None if is_winner else turn_elim,
                    eliminated_by_id=None if is_winner else elim_by,
                    victory_condition=vcond if is_winner else None,
                    loss_condition=lcond if not is_winner else None,
                    notes=notes,
                )
                db.session.add(entry)

            db.session.commit()
            _sync_excel()
            flash('Partita aggiunta con successo!', 'success')
            return redirect(url_for('main.dashboard'))

        except Exception as e:
            db.session.rollback()
            flash(f'Errore: {e}', 'error')

    return render_template(
        'game_form.html',
        players=players,
        commanders=commanders,
        win_conditions=WIN_CONDITIONS,
        loss_conditions=LOSS_CONDITIONS,
        today=date.today().isoformat(),
        game=None,
    )


@main.route('/game/<int:game_id>/edit', methods=['GET', 'POST'])
def game_edit(game_id):
    game = Game.query.get_or_404(game_id)
    players = Player.query.order_by(Player.name).all()
    commanders = Commander.query.order_by(Commander.name).all()

    if request.method == 'POST':
        try:
            game.date = datetime.strptime(request.form['date'], '%Y-%m-%d').date()
            game.total_turns = request.form.get('total_turns', type=int)
            num_players = request.form.get('num_players', type=int, default=4)

            for entry in game.entries:
                db.session.delete(entry)
            db.session.flush()

            for i in range(num_players):
                player_id = request.form.get(f'p{i}_player', type=int)
                commander_id = request.form.get(f'p{i}_commander', type=int)
                finish = request.form.get(f'p{i}_finish', type=int)
                turn_order = request.form.get(f'p{i}_turn_order', type=int)
                sol_ring = bool(request.form.get(f'p{i}_sol_ring'))
                turn_elim = request.form.get(f'p{i}_turn_eliminated', type=int)
                elim_by = request.form.get(f'p{i}_eliminated_by', type=int)
                vcond = request.form.get(f'p{i}_victory_condition', '').strip() or None
                lcond = request.form.get(f'p{i}_loss_condition', '').strip() or None
                notes = request.form.get(f'p{i}_notes', '').strip() or None
                is_winner = (finish == 1)
                entry = GameEntry(
                    game_id=game.id,
                    player_id=player_id,
                    commander_id=commander_id,
                    finish=finish,
                    turn_order=turn_order,
                    sol_ring_t1=sol_ring,
                    turn_eliminated=None if is_winner else turn_elim,
                    eliminated_by_id=None if is_winner else elim_by,
                    victory_condition=vcond if is_winner else None,
                    loss_condition=lcond if not is_winner else None,
                    notes=notes,
                )
                db.session.add(entry)

            db.session.commit()
            _sync_excel()
            flash('Partita aggiornata!', 'success')
            return redirect(url_for('main.history'))

        except Exception as e:
            db.session.rollback()
            flash(f'Errore: {e}', 'error')

    return render_template(
        'game_form.html',
        players=players,
        commanders=commanders,
        win_conditions=WIN_CONDITIONS,
        loss_conditions=LOSS_CONDITIONS,
        today=date.today().isoformat(),
        game=game,
    )


@main.route('/game/<int:game_id>/delete', methods=['POST'])
def game_delete(game_id):
    game = Game.query.get_or_404(game_id)
    db.session.delete(game)
    db.session.commit()
    _sync_excel()
    flash('Partita eliminata.', 'success')
    return redirect(url_for('main.history'))


# ── history ───────────────────────────────────────────────────────────────────

@main.route('/history')
def history():
    page = request.args.get('page', 1, type=int)
    player_filter = request.args.get('player', '')
    commander_filter = request.args.get('commander', '')

    query = Game.query.order_by(Game.date.desc())

    if player_filter:
        p = Player.query.filter_by(name=player_filter).first()
        if p:
            ids = [e.game_id for e in GameEntry.query.filter_by(player_id=p.id).all()]
            query = query.filter(Game.id.in_(ids))

    if commander_filter:
        c = Commander.query.filter_by(name=commander_filter).first()
        if c:
            ids = [e.game_id for e in GameEntry.query.filter_by(commander_id=c.id).all()]
            query = query.filter(Game.id.in_(ids))

    games = query.paginate(page=page, per_page=15, error_out=False)
    players = Player.query.order_by(Player.name).all()
    commanders = Commander.query.order_by(Commander.name).all()

    return render_template(
        'history.html',
        games=games,
        players=players,
        commanders=commanders,
        player_filter=player_filter,
        commander_filter=commander_filter,
    )


# ── commanders ────────────────────────────────────────────────────────────────

@main.route('/commanders')
def commanders():
    cmds = Commander.query.filter(
        Commander.name.isnot(None),
        Commander.name != '',
        Commander.name.regexp_match(r'[A-Za-zÀ-ÿ\']'),
    ).order_by(Commander.name).all()
    stats = []
    for c in cmds:
        entries = GameEntry.query.filter_by(commander_id=c.id).all()
        gp = len(entries)
        wins = sum(1 for e in entries if e.finish == 1)
        kills = GameEntry.query.filter_by(eliminated_by_id=c.id).count()
        stats.append({
            'commander': c,
            'games': gp,
            'wins': wins,
            'kills': kills,
            'win_rate': round(wins / gp * 100, 1) if gp else 0,
        })
    from app.models import COLOR_NAMES
    return render_template('commanders.html', stats=stats, color_names=COLOR_NAMES)


@main.route('/commanders/add', methods=['POST'])
def commander_add():
    name = request.form.get('name', '').strip()
    color = request.form.get('color_identity', '').strip().upper()
    if name and color:
        if not Commander.query.filter_by(name=name).first():
            db.session.add(Commander(name=name, color_identity=color))
            db.session.commit()
            flash(f'Commander "{name}" aggiunto.', 'success')
        else:
            flash(f'Commander "{name}" esiste già.', 'error')
    return redirect(url_for('main.commanders'))


@main.route('/commanders/<int:cmd_id>/edit', methods=['POST'])
def commander_edit(cmd_id):
    c = Commander.query.get_or_404(cmd_id)
    c.name = request.form.get('name', c.name).strip()
    c.color_identity = request.form.get('color_identity', c.color_identity).strip().upper()
    db.session.commit()
    flash('Commander aggiornato.', 'success')
    return redirect(url_for('main.commanders'))


@main.route('/commanders/<int:cmd_id>/delete', methods=['POST'])
def commander_delete(cmd_id):
    c = Commander.query.get_or_404(cmd_id)
    if GameEntry.query.filter_by(commander_id=cmd_id).count() > 0:
        flash('Impossibile eliminare: il commander ha partite registrate.', 'error')
    else:
        db.session.delete(c)
        db.session.commit()
        flash('Commander eliminato.', 'success')
    return redirect(url_for('main.commanders'))


# ── import / export ───────────────────────────────────────────────────────────

@main.route('/import', methods=['GET', 'POST'])
def import_excel():
    if request.method == 'POST':
        f = request.files.get('file')
        if f and f.filename.endswith('.xlsx'):
            import tempfile
            with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                f.save(tmp.name)
                tmp_path = tmp.name
            try:
                from app.import_excel import import_from_excel
                count = import_from_excel(tmp_path)
                _sync_excel()
                flash(f'Importate {count} partite con successo!', 'success')
            except Exception as e:
                flash(f'Importazione fallita: {e}', 'error')
            finally:
                os.unlink(tmp_path)
        else:
            flash('Seleziona un file .xlsx valido.', 'error')
        return redirect(url_for('main.dashboard'))
    return render_template('import.html')


@main.route('/export')
def export_excel():
    from app.excel_sync import sync_to_excel
    path = sync_to_excel(current_app.config['DATA_DIR'])
    return send_file(path, as_attachment=True, download_name='mtg_tracking.xlsx')
