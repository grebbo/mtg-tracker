import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import Game

HEADERS = [
    'Date mm/dd/yyyy', 'Deck Archtype', 'Pilot', 'Finish', 'Turn Order',
    'Turn 1 Sol Ring', 'Turns', 'Turn Eliminated', 'Eliminated By',
    'Victory Condition', 'Loss Condition', 'Notes',
]
COL_WIDTHS = [16, 18, 12, 8, 12, 15, 8, 16, 18, 18, 18, 25]


def sync_to_excel(data_dir: str) -> str:
    path = os.path.join(data_dir, 'mtg_tracking.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tracking Sheet'

    # ── header row ────────────────────────────────────────────────────────────
    hdr_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    hdr_font = Font(bold=True, color='FFFFFF', size=10)
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30

    # ── data rows ─────────────────────────────────────────────────────────────
    row = 2
    alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    game_idx = 0

    games = Game.query.order_by(Game.date).all()
    for game in games:
        entries = sorted(game.entries, key=lambda e: (e.finish or 99, e.turn_order or 99))
        fill = alt_fill if game_idx % 2 == 0 else None
        first = True

        for entry in entries:
            for col in range(1, 13):
                c = ws.cell(row=row, column=col)
                if fill:
                    c.fill = fill
                c.alignment = Alignment(vertical='center')

            if first:
                ws.cell(row=row, column=1, value=game.date.strftime('%m/%d/%Y'))
                if game.total_turns:
                    ws.cell(row=row, column=7, value=game.total_turns)

            ws.cell(row=row, column=2, value=entry.commander.name if entry.commander else '')
            ws.cell(row=row, column=3, value=entry.player.name if entry.player else '')
            ws.cell(row=row, column=4, value=entry.finish)
            ws.cell(row=row, column=5, value=entry.turn_order)

            if entry.sol_ring_t1:
                ws.cell(row=row, column=6, value='Yes')

            if entry.turn_eliminated:
                ws.cell(row=row, column=8, value=entry.turn_eliminated)

            if entry.eliminated_by_commander:
                ws.cell(row=row, column=9, value=entry.eliminated_by_commander.name)

            if entry.victory_condition:
                ws.cell(row=row, column=10, value=entry.victory_condition)

            if entry.loss_condition:
                ws.cell(row=row, column=11, value=entry.loss_condition)

            if entry.notes:
                ws.cell(row=row, column=12, value=entry.notes)

            first = False
            row += 1

        row += 1  # blank row between games
        game_idx += 1

    # ── column widths ─────────────────────────────────────────────────────────
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws.freeze_panes = 'A2'
    wb.save(path)
    return path
