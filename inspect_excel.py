"""
Run: python inspect_excel.py <percorso_del_tuo_excel.xlsx>
"""
import sys
from openpyxl import load_workbook

if len(sys.argv) < 2:
    print("Uso: python inspect_excel.py <file.xlsx>")
    sys.exit(1)

path = sys.argv[1]
wb = load_workbook(path, data_only=True)

print(f"\nFogli presenti: {wb.sheetnames}\n")

# Use first sheet or one with 'track' in name
ws = wb.worksheets[0]
for sheet in wb.worksheets:
    if 'track' in sheet.title.lower():
        ws = sheet
        break

print(f"Foglio selezionato: '{ws.title}'\n")

# Print first 10 rows with column indices
print("=" * 80)
for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
    print(f"\nRiga {r_idx}:")
    for c_idx, val in enumerate(row[:16]):
        if val is not None and str(val).strip():
            print(f"  col[{c_idx:2d}] = {repr(val)}")
print("=" * 80)
